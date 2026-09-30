import base64
import re
from openai import OpenAI
from pypdf import PdfReader
from pathlib import Path
import math
from io import BytesIO
from functools import lru_cache
from docx import Document
from docx.shared import Pt, Cm, Emu, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from openpyxl.styles import Font
from openpyxl.cell.text import InlineFont
from openpyxl.cell.rich_text import CellRichText, TextBlock


# ==========================================================================================
# Functions
# ==========================================================================================
def load_file_base64(path: str) -> str:
    with open(path, "rb") as f:
        return base64.b64encode(f.read()).decode()

def extract_text_from_pdf(uploaded_file) -> str:
    reader = PdfReader(uploaded_file)
    full_text = ""
    for page in reader.pages:
        text = page.extract_text()
        if text:
            full_text += text + "\n"
    return full_text

SUBSCORE_KEYS = {"tech": "TECH", "exp": "EXP", "edu": "EDU", "proj": "PROJ"}


def extract_subscores(text: str) -> dict:
    """زیرنمره‌ی (۰ تا ۱۰۰) هر معیار را از خروجی مدل می‌خواند؛ معیاری که پیدا نشود در خروجی نمی‌آید."""
    if not text:
        return {}
    for digits in ("۰۱۲۳۴۵۶۷۸۹", "٠١٢٣٤٥٦٧٨٩"):
        for p, e in zip(digits, "0123456789"):
            text = text.replace(p, e)
    subs = {}
    for key, label in SUBSCORE_KEYS.items():
        m = re.search(rf"(?<![A-Za-z]){label}[\s*_]*[:：][\s*_]*(100|[1-9]?\d)(?!\d)", text, re.IGNORECASE)
        if m:
            subs[key] = int(m.group(1))
    return subs


def extract_candidate_name(text: str) -> str:
    """نام متقاضی (به فارسی) را از خط NAME در خروجی مدل می‌خواند."""
    if not text:
        return "-"
    m = re.search(r"(?im)^\s*[*_]*NAME[*_\s]*[:：]\s*(.+?)\s*$", text)
    if not m:
        return "-"
    name = re.sub(r"[*_`]+", "", m.group(1)).strip(" \t:：")
    if not name or "نامشخص" in name or name.lower() in {"unknown", "n/a", "na", "-"}:
        return "-"
    return name


def extract_score(text: str, weights: dict = None):
    """
    نمره‌ی نهایی = Σ(ضریب × زیرنمره) ÷ Σ(ضرایب)  ← محاسبه در پایتون (نه توسط مدل).
    ضرایب به‌صورت خودکار نرمال می‌شوند؛ معیاری که ضریبش صفر است در محاسبه نمی‌آید.
    اگر زیرنمره‌ی یکی از معیارهای دارای ضریب در خروجی مدل نباشد، "-" برمی‌گردد.
    """
    subs = extract_subscores(text)
    if weights is None:
        weights = {k: 1 for k in SUBSCORE_KEYS}
    used = {k: w for k, w in weights.items() if k in SUBSCORE_KEYS and w > 0}
    if not subs or not used or any(k not in subs for k in used):
        return "-"
    weighted = sum(w * subs[k] for k, w in used.items()) / sum(used.values())
    return str(int(weighted + 0.5))

# ==========================================================================================
# Resume Evaluation Bot
# ==========================================================================================
class ResumeEvaluatorBot:
    def __init__(self, api_key: str, base_url: str):
        self.client = OpenAI(api_key=api_key, base_url=base_url)

    def build_system_prompt(self, weights: dict) -> str:
        # ضرایب دیگر به مدل داده نمی‌شوند؛ وزن‌دهی نهایی در extract_score و در پایتون انجام می‌شود.
        return """
You are a professional AI-based recruitment and competency assessment expert.

Your task is to evaluate a candidate's resume strictly and only based on the provided job competencies.

Evaluation rules:
1. You must NOT use any hidden assumptions about the job.
2. You must NOT introduce additional competencies.
3. You must rely only on the resume content and the provided competencies.
4. If a competency is not evidenced in the resume, clearly state it is not demonstrated.

Your output must be written in Persian.

Your response must follow exactly this structure:

Paragraph 1:
A short overall evaluation (about 4 to 6 sentences) of the candidate's suitability for the role.
This paragraph must explicitly refer to the candidate's previous work experience, emphasizing the projects
the candidate has participated in and the skills the candidate has acquired through them.
If the resume does not demonstrate such experience, projects or skills, clearly state that.

Paragraph 2:
A competency-by-competency assessment in bullet points.

Paragraph 3:
Exactly five lines in this exact machine-readable format, nothing else on these lines:
NAME: <candidate's full name, written in Persian script, regardless of the language used in the resume>
TECH: <score for technical skills, an integer from 0 to 100>
EXP: <score for work experience, an integer from 0 to 100>
EDU: <score for education, an integer from 0 to 100>
PROJ: <score for projects, an integer from 0 to 100>
If the candidate's name cannot be determined from the resume with reasonable confidence, write "NAME: نامشخص".
Each score must independently reflect how well the resume demonstrates the provided competencies
within that criterion only. If a criterion is not evidenced in the resume, give it a low score.
Do NOT compute or mention any overall score.

Paragraph 4:
A section titled "نقاط قوت رزومه" listing the strongest points.

Paragraph 5:
A section titled "نقاط ضعف رزومه" listing the weakest areas.

Do not use any headings or titles other than "نقاط قوت رزومه" and "نقاط ضعف رزومه".
Do not mention that you are an AI model.
Be precise, professional and critical.
"""

    def evaluate(self, resume_text: str, competencies: str, weights: dict) -> str:
        messages = [
            {"role": "system", "content": self.build_system_prompt(weights)},
            {
                "role": "user",
                "content": (
                    "شایستگی‌های شغلی مدنظر:\n\n"
                    f"{competencies}\n\n"
                    "-------------------------\n\n"
                    "متن استخراج شده از رزومه:\n\n"
                    f"{resume_text}"
                )
            }
        ]
        response = self.client.chat.completions.create(
            model= "gpt-6-luna",    # gpt-6-luna
            messages=messages,
            temperature=0.1,
            max_tokens=2000     # 1200  # 2000
        )
        return response.choices[0].message.content

# ==========================================================================================
# Assets
# ==========================================================================================
font_base64 = load_file_base64(Path.cwd() / "Assets/BKoodkBd.ttf")
logo_base64 = load_file_base64(Path.cwd() / "Assets/Monenco Iran Logo.png")

# ==========================================================================================
# Word Report (docx) - گزارش تحلیلی رزومه‌ها
# ==========================================================================================
REPORT_TEMPLATE_PATH = Path.cwd() / "Assets/Report_Template.docx"   # <-- نام فایل تمپلیت خام گزارش
REPORT_FONT_PATH     = Path.cwd() / "Assets/BKoodkBd.ttf"           # همان فونت سامانه
REPORT_FONT_SIZE_PT  = 13                                           # اندازه فونت گزارش
REPORT_LINE_PT       = 22                                           # ارتفاع ثابت هر خط (نقطه)
REPORT_GAP_PT        = 4                                            # فاصله بعد از هر پاراگراف (نقطه)
REPORT_SLACK_PT      = 40                                           # حاشیه اطمینان ارتفاع صفحه (نقطه)
REPORT_CHAR_PT       = 7.2                                          # تخمین محافظه‌کارانه عرض هر نویسه (نقطه)
REPORT_HEAD_COLOR    = RGBColor(0x1F, 0x38, 0x64)                   # سرمه‌ای (هم‌رنگ فیلدهای فرانت)
REPORT_GREEN_COLOR   = RGBColor(0x1E, 0x8E, 0x3E)                   # سبز: نقاط قوت
REPORT_RED_COLOR     = RGBColor(0xC0, 0x39, 0x2B)                   # قرمز: نقاط ضعف

_PPR_AFTER_BIDI = (
    "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing",
    "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
    "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)


@lru_cache(maxsize=1)
def _report_font():
    """نام خانواده و Bold بودن فونت را مستقیم از فایل ttf می‌خواند (عین فونت سامانه)."""
    try:
        from fontTools.ttLib import TTFont
        names = TTFont(str(REPORT_FONT_PATH))["name"]
        family = names.getDebugName(1) or "B Koodak"
        return family, "bold" in (names.getDebugName(2) or "").lower()
    except Exception:
        return "B Koodak", True


def get_score_level(score: int) -> str:
    """سطح‌بندی چهارگانه‌ی نمره (نمره ۶۰ در سطح «معمولی» محاسبه می‌شود)."""
    if score < 60:
        return "نامناسب برای موقعیت شغلی"
    if score <= 70:
        return "معمولی - اولویت با سایر گزینه‌ها"
    if score <= 85:
        return "قابل قبول"
    return "مناسب - اولویت برای مصاحبه"


def _style_run(run, color=None):
    family, is_bold = _report_font()
    run.font.name = family
    run.font.size = Pt(REPORT_FONT_SIZE_PT)
    run.font.bold = is_bold
    if color is not None:
        run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), family)
    if is_bold:
        rPr.find(qn("w:b")).addnext(OxmlElement("w:bCs"))
    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), str(REPORT_FONT_SIZE_PT * 2))
    rPr.append(szCs)
    rPr.append(OxmlElement("w:rtl"))


def _add_par(container, text="", color=None, bullet=False, first=None, justify=False):
    """پاراگراف راست‌به‌چپ با ارتفاع خط ثابت؛ ارتفاع ثابت باعث می‌شود جانمایی صفحه قابل پیش‌بینی باشد."""
    p = first if first is not None else container.add_paragraph()
    p._p.get_or_add_pPr().insert_element_before(OxmlElement("w:bidi"), *_PPR_AFTER_BIDI)
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(REPORT_GAP_PT)
    pf.line_spacing = Pt(REPORT_LINE_PT)
    if justify:
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bullet:
        pf.left_indent = Cm(0.6)
        pf.first_line_indent = Cm(-0.6)
        text = "•\t" + text
    if text:
        _style_run(p.add_run(text), color)
    return p


def _tiny_par(doc, page_break=False):
    """پاراگراف ۱ نقطه‌ای: جداکننده‌ی صفحه‌ها و پاراگراف پایانی الزامی بعد از جدول."""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = Pt(1)
    pf.page_break_before = page_break
    return p


def _clean_line(s: str) -> str:
    s = re.sub(r"[*#`_]+", "", s)
    s = re.sub(r"^\s*(?:[-–—•●▪]+|\d+[.)]|[۰-۹]+[.)])\s*", "", s)
    return s.strip(" \t:")


def _split_items(block: str):
    return [c for c in (_clean_line(l) for l in block.splitlines()) if c]


def _parse_resume_analysis(text: str):
    """متن خروجی مدل را به (ارزیابی کلی، نقاط قوت، نقاط ضعف) تبدیل می‌کند."""
    text = text or ""
    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    overall = " ".join(_split_items(paragraphs[0])) if paragraphs else ""
    m_s, m_w = "نقاط قوت رزومه", "نقاط ضعف رزومه"
    i_s = text.find(m_s)
    i_w = text.find(m_w, i_s + 1 if i_s != -1 else 0)
    if i_s != -1 and i_w != -1:
        return overall, _split_items(text[i_s + len(m_s):i_w]), _split_items(text[i_w + len(m_w):])
    if i_s != -1:
        return overall, _split_items(text[i_s + len(m_s):]), []
    if i_w != -1:
        return overall, [], _split_items(text[i_w + len(m_w):])
    return " ".join(_split_items(re.sub(r"(?im)^\s*[*_]*(NAME|TECH|EXP|EDU|PROJ)\b.*$", "", text))), [], []


def _truncate(text: str, max_chars: int) -> str:
    if len(text) <= max_chars:
        return text
    cut = text[:max(max_chars - 1, 1)].rsplit(" ", 1)[0].rstrip("،,.;:- ")
    return cut + "…"


def _fit_items(items, cpl, budget_pt, indent=0):
    """آیتم‌ها را تا سقف ارتفاع مجاز جا می‌دهد؛ آیتم آخر در صورت نیاز کوتاه می‌شود."""
    eff = max(cpl - indent, 1)
    fitted, used = [], 0
    for it in items:
        cost = math.ceil(len(it) / eff) * REPORT_LINE_PT + REPORT_GAP_PT
        if used + cost <= budget_pt:
            fitted.append(it)
            used += cost
            continue
        room = int((budget_pt - used - REPORT_GAP_PT) // REPORT_LINE_PT)
        if room >= 1:
            t = _truncate(it, room * eff)
            fitted.append(t)
            used += math.ceil(len(t) / eff) * REPORT_LINE_PT + REPORT_GAP_PT
        break
    return fitted, used


def build_resume_report(ranked_results) -> BytesIO:
    """
    گزارش تحلیلی Word را از روی تمپلیت خام می‌سازد.
    ranked_results: فهرست نتایج به ترتیب رتبه (همان ترتیب فایل Excel).
    هر رزومه دقیقاً یک صفحه: یک جدول تک‌سلولی با ارتفاع «دقیق» (exact) که هم از سرریز به صفحه بعد
    جلوگیری می‌کند و هم با شکست صفحه‌ی اجباری، هر رزومه را در صفحه‌ی مستقل قرار می‌دهد.
    """
    doc = Document(str(REPORT_TEMPLATE_PATH))

    # اگر بدنه‌ی تمپلیت خالی است، پاراگراف‌های خالی را حذف کن تا صفحه‌ی اول خالی نماند
    has_content = bool(doc.tables) or any(
        p.text.strip() or p._p.xpath(".//w:drawing | .//w:pict | .//w:sectPr | .//w:br")
        for p in doc.paragraphs
    )
    if not has_content:
        for p in list(doc.paragraphs):
            p._element.getparent().remove(p._element)

    sec = doc.sections[-1]
    text_w = Emu(sec.page_width - sec.left_margin - sec.right_margin)
    page_h_pt = Emu(sec.page_height - sec.top_margin - sec.bottom_margin).pt - REPORT_SLACK_PT
    cpl = int((text_w.pt - 30) / REPORT_CHAR_PT)
    content_pt = page_h_pt - 12

    for idx, r in enumerate(ranked_results, start=1):
        if idx > 1 or has_content:
            _tiny_par(doc, page_break=True)

        table = doc.add_table(rows=1, cols=1)
        table.autofit = False
        table.columns[0].width = text_w
        row = table.rows[0]
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
        h = OxmlElement("w:trHeight")
        h.set(qn("w:val"), str(int(page_h_pt * 20)))
        h.set(qn("w:hRule"), "exact")
        trPr.append(h)
        cell = row.cells[0]
        cell.width = text_w

        overall, strengths, weaknesses = _parse_resume_analysis(r.get("result_text", ""))
        name_line = f"نام فایل رزومه: {r['file_name']}"
        applicant_line = f"نام متقاضی: {r.get('applicant_name', '-')}"
        level_line = f"سطح ارزیابی: {get_score_level(r['score_raw'])}"
        head_lines = [f"رتبه: {idx}", name_line, applicant_line, f"نمره ارزیابی: {r['score_raw']} از 100", level_line]
        head_pt = sum(math.ceil(len(t) / cpl) * REPORT_LINE_PT + REPORT_GAP_PT for t in head_lines)
        heading_pt = 3 * (REPORT_LINE_PT + REPORT_GAP_PT)
        remaining = content_pt - head_pt - heading_pt

        ov, used_o = _fit_items([overall] if overall else [], cpl, remaining * 0.40)
        rem = remaining - used_o
        st_items, used_s = _fit_items(strengths, cpl, rem / 2, indent=3)
        wk_items, _ = _fit_items(weaknesses, cpl, rem - used_s, indent=3)

        _add_par(cell, head_lines[0], color=REPORT_HEAD_COLOR, first=cell.paragraphs[0])
        _add_par(cell, head_lines[1])
        _add_par(cell, head_lines[2])
        _add_par(cell, head_lines[3], color=REPORT_HEAD_COLOR)
        _add_par(cell, head_lines[4], color=REPORT_HEAD_COLOR)

        _add_par(cell, "تحلیل کلی رزومه", color=REPORT_HEAD_COLOR)
        _add_par(cell, ov[0] if ov else "موردی ثبت نشده است.", justify=True)
        _add_par(cell, "نقاط قوت رزومه", color=REPORT_GREEN_COLOR)
        for t in (st_items or ["موردی ثبت نشده است."]):
            _add_par(cell, t, color=REPORT_GREEN_COLOR, bullet=True)
        _add_par(cell, "نقاط ضعف رزومه", color=REPORT_RED_COLOR)
        for t in (wk_items or ["موردی ثبت نشده است."]):
            _add_par(cell, t, color=REPORT_RED_COLOR, bullet=True)

    _tiny_par(doc)   # پاراگراف پایانی الزامی پس از آخرین جدول

    out = BytesIO()
    doc.save(out)
    out.seek(0)
    return out


# ==========================================================================================
# Excel Styling - فونت فارسی هم‌فونت گزارش، فونت انگلیسی Cambria
# ==========================================================================================
EXCEL_LATIN_FONT = "Cambria"
EXCEL_FONT_SIZE  = 11
_PERSIAN_CHAR_RE = re.compile(r"[\u0600-\u06FF\u0750-\u077F\uFB50-\uFDFF\uFE70-\uFEFF]")


def _excel_runs(text: str):
    """رشته را به قطعه‌های (متن، فارسی؟) تقسیم می‌کند؛ نویسه‌های خنثی (فاصله، نقطه، خط‌تیره...) به قطعه‌ی مجاور می‌چسبند."""
    kinds = []
    for ch in text:
        if _PERSIAN_CHAR_RE.match(ch):
            kinds.append(True)
        elif ch.isalnum():
            kinds.append(False)
        else:
            kinds.append(None)
    last = None
    for i, k in enumerate(kinds):          # پر کردن خنثی‌ها با نوع قبلی
        if k is None:
            kinds[i] = last
        else:
            last = k
    nxt = None
    for i in range(len(kinds) - 1, -1, -1):  # خنثی‌های ابتدای رشته با نوع بعدی
        if kinds[i] is None:
            kinds[i] = nxt
        else:
            nxt = kinds[i]
    runs = []
    for ch, k in zip(text, kinds):
        k = bool(k)
        if runs and runs[-1][1] == k:
            runs[-1][0] += ch
        else:
            runs.append([ch, k])
    return [(t, k) for t, k in runs]


def style_excel_sheet(ws):
    """
    فونت همه‌ی سلول‌های شیت را تنظیم می‌کند:
    متن فارسی = فونت گزارش تحلیلی (BKoodkBd)، متن و اعداد انگلیسی = Cambria.
    اگر یک سلول ترکیبی باشد (مثلاً نام فایل فارسی + انگلیسی)، هر بخش با فونت خودش نوشته می‌شود.
    """
    family, is_bold = _report_font()
    fa_font = Font(name=family, size=EXCEL_FONT_SIZE, bold=is_bold)
    en_font = Font(name=EXCEL_LATIN_FONT, size=EXCEL_FONT_SIZE)
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if v is None:
                continue
            if not isinstance(v, str):          # اعداد (رتبه، نمره) => Cambria
                cell.font = en_font
                continue
            runs = _excel_runs(v)
            if not runs:
                continue
            kinds = {k for _, k in runs}
            if kinds == {True}:
                cell.font = fa_font
            elif kinds == {False}:
                cell.font = en_font
            else:
                cell.font = en_font
                cell.value = CellRichText(*[
                    TextBlock(
                        InlineFont(rFont=(family if k else EXCEL_LATIN_FONT), sz=EXCEL_FONT_SIZE,
                                   b=(True if (k and is_bold) else None)),
                        t)
                    for t, k in runs
                ])
