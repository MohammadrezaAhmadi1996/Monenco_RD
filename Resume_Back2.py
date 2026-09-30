import base64
import re
import json
from openai import OpenAI
from pypdf import PdfReader
from pathlib import Path
from io import BytesIO
from functools import lru_cache
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from openpyxl.styles import Font
from openpyxl.cell.text import InlineFont
from openpyxl.cell.rich_text import CellRichText, TextBlock


# ==========================================================================================
# Functions
# ==========================================================================================
def load_file_base64(path: str) -> str:
    with open(path, "rb") as f:
        return base64.b64encode(f.read()).decode()

def extract_text_from_pdf(uploaded_file) -> str:
    reader = PdfReader(uploaded_file)
    full_text = ""
    for page in reader.pages:
        text = page.extract_text()
        if text:
            full_text += text + "\n"
    return full_text

SUBSCORE_KEYS = {"tech": "TECH", "exp": "EXP", "edu": "EDU", "proj": "PROJ"}


def extract_subscores(text: str) -> dict:
    """زیرنمره‌ی (۰ تا ۱۰۰) هر معیار را از خروجی مدل می‌خواند؛ معیاری که پیدا نشود در خروجی نمی‌آید."""
    if not text:
        return {}
    for digits in ("۰۱۲۳۴۵۶۷۸۹", "٠١٢٣٤٥٦٧٨٩"):
        for p, e in zip(digits, "0123456789"):
            text = text.replace(p, e)
    subs = {}
    for key, label in SUBSCORE_KEYS.items():
        m = re.search(rf"(?<![A-Za-z]){label}[\s*_]*[:：][\s*_]*(100|[1-9]?\d)(?!\d)", text, re.IGNORECASE)
        if m:
            subs[key] = int(m.group(1))
    return subs


def extract_candidate_name(text: str) -> str:
    """نام متقاضی (به فارسی) را از خط NAME در خروجی مدل می‌خواند."""
    if not text:
        return "-"
    m = re.search(r"(?im)^\s*[*_]*NAME[*_\s]*[:：]\s*(.+?)\s*$", text)
    if not m:
        return "-"
    name = re.sub(r"[*_`]+", "", m.group(1)).strip(" \t:：")
    if not name or "نامشخص" in name or name.lower() in {"unknown", "n/a", "na", "-"}:
        return "-"
    return name


def extract_score(text: str, weights: dict = None):
    """
    نمره‌ی نهایی = Σ(ضریب × زیرنمره) ÷ Σ(ضرایب)  ← محاسبه در پایتون (نه توسط مدل).
    ضرایب به‌صورت خودکار نرمال می‌شوند؛ معیاری که ضریبش صفر است در محاسبه نمی‌آید.
    اگر زیرنمره‌ی یکی از معیارهای دارای ضریب در خروجی مدل نباشد، "-" برمی‌گردد.
    """
    subs = extract_subscores(text)
    if weights is None:
        weights = {k: 1 for k in SUBSCORE_KEYS}
    used = {k: w for k, w in weights.items() if k in SUBSCORE_KEYS and w > 0}
    if not subs or not used or any(k not in subs for k in used):
        return "-"
    weighted = sum(w * subs[k] for k, w in used.items()) / sum(used.values())
    return str(int(weighted + 0.5))

# ==========================================================================================
# Resume Evaluation Bot
# ==========================================================================================
class ResumeEvaluatorBot:
    def __init__(self, api_key: str, base_url: str):
        self.client = OpenAI(api_key=api_key, base_url=base_url)

    def build_system_prompt(self, weights: dict) -> str:
        # ضرایب دیگر به مدل داده نمی‌شوند؛ وزن‌دهی نهایی در extract_score و در پایتون انجام می‌شود.
        return """
You are a professional AI-based recruitment and competency assessment expert.

Your task is to evaluate a candidate's resume strictly and only based on the provided job competencies.

Evaluation rules:
1. You must NOT use any hidden assumptions about the job.
2. You must NOT introduce additional competencies.
3. You must rely only on the resume content and the provided competencies.
4. If a competency is not evidenced in the resume, clearly state it is not demonstrated.
5. Every strength and weakness you report must be judged RELATIVE to this specific job position and the
   provided competencies, never in absolute or general terms. A fact from the resume that would be impressive
   or concerning in general, but is unrelated to this job position and these competencies, must NOT be reported
   as a strength or a weakness (for example, an outstanding GPA in a field unrelated to a software-engineering
   position is not a strength for that position, even though a high GPA is impressive in general).

Your output must be written in Persian.

Your response must follow exactly this structure:

Paragraph 1:
A thorough and detailed overall evaluation of the candidate's suitability for the role. There is no length
limit: write as many sentences as needed to fully and completely cover the candidate's relevant educational
background, technical skills and work history, so that a reader would not need to go back to the original
resume file to learn any important, job-relevant point. This paragraph must explicitly discuss the candidate's
previous work experience, emphasizing the projects the candidate has participated in and the skills the
candidate has acquired through them. If the resume does not demonstrate such experience, projects or skills,
clearly state that. Do not include facts that are not relevant to this job position or the provided
competencies (see rule 5 above).

Paragraph 2:
A competency-by-competency assessment in bullet points.

Paragraph 3:
Exactly five lines in this exact machine-readable format, nothing else on these lines:
NAME: <candidate's full name, written in Persian script, regardless of the language used in the resume>
TECH: <score for technical skills, an integer from 0 to 100>
EXP: <score for work experience, an integer from 0 to 100>
EDU: <score for education, an integer from 0 to 100>
PROJ: <score for projects, an integer from 0 to 100>
If the candidate's name cannot be determined from the resume with reasonable confidence, write "NAME: نامشخص".
Each score must independently reflect how well the resume demonstrates the provided competencies
within that criterion only. If a criterion is not evidenced in the resume, give it a low score.
Do NOT compute or mention any overall score.

Paragraph 4:
A section titled "نقاط قوت رزومه" listing the strongest points. List a point here only if it is genuinely
relevant to this job position and the provided competencies (see rule 5 above); omit anything impressive
that is unrelated to them.

Paragraph 5:
A section titled "نقاط ضعف رزومه" listing the weakest areas, identified the same way: only weaknesses that
matter for this specific job position and the provided competencies (see rule 5 above).

Do not use any headings or titles other than "نقاط قوت رزومه" and "نقاط ضعف رزومه".
Do not mention that you are an AI model.
Be precise, professional and critical.
"""

    def evaluate(self, resume_text: str, competencies: str, weights: dict) -> str:
        messages = [
            {"role": "system", "content": self.build_system_prompt(weights)},
            {
                "role": "user",
                "content": (
                    "شایستگی‌های شغلی مدنظر:\n\n"
                    f"{competencies}\n\n"
                    "-------------------------\n\n"
                    "متن استخراج شده از رزومه:\n\n"
                    f"{resume_text}"
                )
            }
        ]
        response = self.client.chat.completions.create(
            model= "gpt-6-luna",    # gpt-6-luna
            messages=messages,
            temperature=0.1,
            max_tokens=2000     # 1200  # 2000
        )
        return response.choices[0].message.content



# ==========================================================================================
# Interview Question Generator - Phase 2
# ==========================================================================================
class InterviewQuestionGenerator:
    """
    تولید سوالات مصاحبه در دو سطح:
      1) سوالات عمومی مشترک برای همه متقاضیان منتخب
      2) سوالات اختصاصی برای هر رزومه منتخب
    سوالات دسته‌های 3 و 4 صرفاً برای راستی‌آزمایی و سنجش شایستگی‌های مرتبط
    با شغل و ادعاهای مستندشده در رزومه طراحی می‌شوند.
    """

    COMMON_CATEGORIES = {
        "general_ai": "۱. دانش عمومی و مفاهیم هوش مصنوعی، یادگیری ماشین، علم داده، تحلیل داده، LLM و مدل‌های زبانی، Prompt Engineering و Agentic AI",
        "behavioral": "۲. سوالات رفتاری، حرفه‌ای و اخلاق کاری برای شناخت شیوه کار، مسئولیت‌پذیری، ارتباط، حل تعارض، یادگیری و صداقت حرفه‌ای",
    }
    PERSONAL_CATEGORIES = {
        "competency": "۳. سوالات تخصصی مبتنی بر شایستگی‌های شغلی اعلام‌شده توسط کارفرما",
        "resume": "۴. سوالات تخصصی مبتنی بر تحصیلات، پروژه‌ها، تجربه‌ها و مهارت‌های صریحاً ذکرشده در رزومه",
    }

    def __init__(self, api_key: str, base_url: str, model: str = "gpt-6-luna"):
        self.client = OpenAI(api_key=api_key, base_url=base_url)
        self.model = model

    @staticmethod
    def _parse_json(content: str) -> dict:
        """JSON را حتی در صورت قرارگرفتن داخل markdown code fence استخراج می‌کند."""
        content = (content or "").strip()
        content = re.sub(r"^```(?:json)?\s*", "", content, flags=re.IGNORECASE)
        content = re.sub(r"\s*```$", "", content)
        try:
            return json.loads(content)
        except json.JSONDecodeError:
            m = re.search(r"\{.*\}", content, flags=re.DOTALL)
            if not m:
                raise ValueError("پاسخ مدل در قالب JSON معتبر دریافت نشد.")
            return json.loads(m.group(0))

    @staticmethod
    def _normalize_questions(data: dict, expected_keys: list[str]) -> dict:
        result = {}
        for key in expected_keys:
            items = data.get(key, [])
            if not isinstance(items, list):
                items = []
            cleaned = []
            for item in items:
                if not isinstance(item, dict):
                    continue
                q = str(item.get("question", "")).strip()
                if not q:
                    continue
                cleaned.append({
                    "question": q,
                    "purpose": str(item.get("purpose", "")).strip(),
                    "assessment_focus": str(item.get("assessment_focus", "")).strip(),
                    "follow_up": str(item.get("follow_up", "")).strip(),
                })
            result[key] = cleaned
        return result

    def _call_json(self, system_prompt: str, user_prompt: str, max_tokens: int = 4500) -> dict:
        response = self.client.chat.completions.create(
            model=self.model,
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": user_prompt},
            ],
            temperature=0.2,
            max_tokens=max_tokens,
        )
        return self._parse_json(response.choices[0].message.content)

    def generate_common_questions(
        self,
        job_title: str,
        competencies: str,
        general_count: int = 10,
        behavioral_count: int = 8,
    ) -> dict:
        system_prompt = f"""
You are a senior technical interviewer and structured interview designer.
Generate a Persian interview question bank for a job candidate in the specified role.

IMPORTANT:
- These questions are COMMON and must be identical for every selected candidate.
- They must NOT depend on any individual resume.
- Category 1 covers general AI/ML/Data/LLM/Prompt Engineering/Agentic AI knowledge.
- Category 2 covers behavioral and professional/ethical work topics.
- Do not ask about protected or highly sensitive personal attributes.
- Do not infer personality from the resume; ask observable, job-relevant behavioral questions instead.
- Questions should be open-ended and suitable for an in-person interview.
- Avoid duplicate questions.
- Return ONLY valid JSON.

JSON schema:
{{
  "general_ai": [
    {{
      "question": "...",
      "purpose": "...",
      "assessment_focus": "...",
      "follow_up": "..."
    }}
  ],
  "behavioral": [
    {{
      "question": "...",
      "purpose": "...",
      "assessment_focus": "...",
      "follow_up": "..."
    }}
  ]
}}
"""
        user_prompt = f"""
عنوان موقعیت شغلی:
{job_title or "ذکر نشده"}

شایستگی‌های شغلی مدنظر کارفرما:
{competencies}

تعداد سوال موردنیاز:
- دسته ۱: {general_count}
- دسته ۲: {behavioral_count}

برای هر سوال، «هدف»، «محور ارزیابی» و یک «سوال پیگیری پیشنهادی» نیز بنویس.
"""
        data = self._call_json(system_prompt, user_prompt, max_tokens=5000)
        result = self._normalize_questions(data, ["general_ai", "behavioral"])
        result["general_ai"] = result["general_ai"][:general_count]
        result["behavioral"] = result["behavioral"][:behavioral_count]
        return result

    def generate_personalized_questions(
        self,
        job_title: str,
        competencies: str,
        resume_text: str,
        general_context: str = "",
        competency_count: int = 8,
        resume_count: int = 8,
    ) -> dict:
        system_prompt = f"""
You are a senior technical interviewer specializing in evidence-based resume verification.

Generate Persian interview questions SPECIFICALLY for the candidate whose resume is supplied.
There are two categories:
3) Questions tied directly to the employer's stated job competencies.
4) Questions tied directly to claims in this candidate's resume: education, projects, work experience,
   technologies, responsibilities, achievements and skills.

CORE RULES:
- Use ONLY evidence explicitly present in the supplied resume and the supplied job competencies.
- Do not invent projects, technologies, employers, degrees, responsibilities or achievements.
- Category 4 must cite the concrete resume claim in the question itself or in its purpose.
- Questions should help the interviewer verify depth, ownership, technical understanding and authenticity.
- Prefer questions that require the candidate to explain architecture, decisions, trade-offs,
  implementation details, metrics, debugging, limitations and lessons learned when those are supported by the resume.
- Category 3 may overlap with category 4 only when the overlap is necessary to test a stated competency.
- Do not ask about protected or highly sensitive personal attributes.
- Do not make a hiring decision or score the candidate.
- Return ONLY valid JSON.

JSON schema:
{{
  "competency": [
    {{
      "question": "...",
      "purpose": "...",
      "assessment_focus": "...",
      "follow_up": "..."
    }}
  ],
  "resume": [
    {{
      "question": "...",
      "purpose": "...",
      "assessment_focus": "...",
      "follow_up": "..."
    }}
  ]
}}
"""
        user_prompt = f"""
عنوان موقعیت شغلی:
{job_title or "ذکر نشده"}

شایستگی‌های شغلی مدنظر کارفرما:
{competencies}

سوالات عمومی تولیدشده برای همه متقاضیان (صرفاً برای جلوگیری از تکرار):
{general_context or "هنوز تولید نشده است."}

-------------------------
متن استخراج‌شده از رزومه این متقاضی:
{resume_text}

تعداد سوال موردنیاز:
- دسته ۳: {competency_count}
- دسته ۴: {resume_count}

سوالات را با اولویت «فنی، علمی، قابل راستی‌آزمایی و مبتنی بر شواهد» طراحی کن.
"""
        data = self._call_json(system_prompt, user_prompt, max_tokens=6000)
        result = self._normalize_questions(data, ["competency", "resume"])
        result["competency"] = result["competency"][:competency_count]
        result["resume"] = result["resume"][:resume_count]
        return result


def flatten_interview_questions(common_questions: dict, personalized_questions: dict, candidate_name: str, file_name: str):
    """ساختار سوالات را برای Excel به ردیف‌های قابل‌فیلتر تبدیل می‌کند."""
    rows = []
    labels = {
        "general_ai": "۱- عمومی هوش مصنوعی و فناوری",
        "behavioral": "۲- رفتاری و حرفه‌ای",
        "competency": "۳- تخصصی شایستگی شغلی",
        "resume": "۴- تخصصی مبتنی بر رزومه",
    }
    for key, items in {**common_questions, **personalized_questions}.items():
        for i, item in enumerate(items, 1):
            rows.append({
                "نام متقاضی": candidate_name,
                "نام فایل رزومه": file_name,
                "دسته سوال": labels.get(key, key),
                "شماره": i,
                "سوال": item.get("question", ""),
                "هدف سوال": item.get("purpose", ""),
                "محور ارزیابی": item.get("assessment_focus", ""),
                "سوال پیگیری": item.get("follow_up", ""),
            })
    return rows


def build_interview_report(candidate_name: str, file_name: str, common_questions: dict, personalized_questions: dict) -> BytesIO:
    """گزارش Word سوالات مصاحبه برای یک متقاضی را تولید می‌کند."""
    doc = Document(str(REPORT_TEMPLATE_PATH))
    has_content = bool(doc.tables) or any(
        p.text.strip() or p._p.xpath(".//w:drawing | .//w:pict | .//w:sectPr | .//w:br")
        for p in doc.paragraphs
    )
    if not has_content:
        for p in list(doc.paragraphs):
            p._element.getparent().remove(p._element)

    _add_par(doc, f"سوالات مصاحبه - {candidate_name}", color=REPORT_HEAD_COLOR, page_break=has_content)
    _add_par(doc, f"نام فایل رزومه: {file_name}")

    sections = [
        ("۱. سوالات عمومی حوزه هوش مصنوعی", common_questions.get("general_ai", [])),
        ("۲. سوالات رفتاری و حرفه‌ای", common_questions.get("behavioral", [])),
        ("۳. سوالات تخصصی شایستگی‌های شغلی", personalized_questions.get("competency", [])),
        ("۴. سوالات تخصصی مبتنی بر رزومه", personalized_questions.get("resume", [])),
    ]
    for title, items in sections:
        _add_par(doc, title, color=REPORT_HEAD_COLOR)
        if not items:
            _add_par(doc, "سوالی تولید نشده است.")
            continue
        for i, item in enumerate(items, 1):
            _add_par(doc, f"{i}. {item.get('question', '')}", justify=True)
            if item.get("purpose"):
                _add_par(doc, f"هدف: {item['purpose']}", color=REPORT_GREEN_COLOR, bullet=True)
            if item.get("assessment_focus"):
                _add_par(doc, f"محور ارزیابی: {item['assessment_focus']}", bullet=True)
            if item.get("follow_up"):
                _add_par(doc, f"پیگیری: {item['follow_up']}", bullet=True)

    out = BytesIO()
    doc.save(out)
    out.seek(0)
    return out


# ==========================================================================================
# Assets
# ==========================================================================================
font_base64 = load_file_base64(Path.cwd() / "Assets/BKoodkBd.ttf")
logo_base64 = load_file_base64(Path.cwd() / "Assets/Monenco Iran Logo.png")

# ==========================================================================================
# Word Report (docx) - گزارش تحلیلی رزومه‌ها
# ==========================================================================================
REPORT_TEMPLATE_PATH = Path.cwd() / "Assets/Report_Template.docx"   # <-- نام فایل تمپلیت خام گزارش
REPORT_FONT_PATH     = Path.cwd() / "Assets/BKoodkBd.ttf"           # همان فونت سامانه
REPORT_FONT_SIZE_PT  = 13                                           # اندازه فونت گزارش
REPORT_LINE_PT       = 22                                           # ارتفاع ثابت هر خط (نقطه)
REPORT_GAP_PT        = 4                                            # فاصله بعد از هر پاراگراف (نقطه)
REPORT_HEAD_COLOR    = RGBColor(0x1F, 0x38, 0x64)                   # سرمه‌ای (هم‌رنگ فیلدهای فرانت)
REPORT_GREEN_COLOR   = RGBColor(0x1E, 0x8E, 0x3E)                   # سبز: نقاط قوت
REPORT_RED_COLOR     = RGBColor(0xC0, 0x39, 0x2B)                   # قرمز: نقاط ضعف

_PPR_AFTER_BIDI = (
    "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing",
    "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
    "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)


@lru_cache(maxsize=1)
def _report_font():
    """نام خانواده و Bold بودن فونت را مستقیم از فایل ttf می‌خواند (عین فونت سامانه)."""
    try:
        from fontTools.ttLib import TTFont
        names = TTFont(str(REPORT_FONT_PATH))["name"]
        family = names.getDebugName(1) or "B Koodak"
        return family, "bold" in (names.getDebugName(2) or "").lower()
    except Exception:
        return "B Koodak", True


def get_score_level(score: int) -> str:
    """سطح‌بندی چهارگانه‌ی نمره (نمره ۶۰ در سطح «معمولی» محاسبه می‌شود)."""
    if score < 60:
        return "نامناسب برای موقعیت شغلی"
    if score <= 70:
        return "معمولی - اولویت با سایر گزینه‌ها"
    if score <= 85:
        return "قابل قبول"
    return "مناسب - اولویت برای مصاحبه"


def _style_run(run, color=None):
    family, is_bold = _report_font()
    run.font.name = family
    run.font.size = Pt(REPORT_FONT_SIZE_PT)
    run.font.bold = is_bold
    if color is not None:
        run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), family)
    if is_bold:
        rPr.find(qn("w:b")).addnext(OxmlElement("w:bCs"))
    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), str(REPORT_FONT_SIZE_PT * 2))
    rPr.append(szCs)
    rPr.append(OxmlElement("w:rtl"))


def _add_par(container, text="", color=None, bullet=False, justify=False, page_break=False):
    """پاراگراف راست‌به‌چپ؛ در صورت page_break=True از ابتدای یک صفحه‌ی جدید شروع می‌شود."""
    p = container.add_paragraph()
    p._p.get_or_add_pPr().insert_element_before(OxmlElement("w:bidi"), *_PPR_AFTER_BIDI)
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(REPORT_GAP_PT)
    pf.line_spacing = Pt(REPORT_LINE_PT)
    pf.page_break_before = page_break
    if justify:
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bullet:
        pf.left_indent = Cm(0.6)
        pf.first_line_indent = Cm(-0.6)
        text = "•\t" + text
    if text:
        _style_run(p.add_run(text), color)
    return p


def _clean_line(s: str) -> str:
    s = re.sub(r"[*#`_]+", "", s)
    s = re.sub(r"^\s*(?:[-–—•●▪]+|\d+[.)]|[۰-۹]+[.)])\s*", "", s)
    return s.strip(" \t:")


def _split_items(block: str):
    return [c for c in (_clean_line(l) for l in block.splitlines()) if c]


def _parse_resume_analysis(text: str):
    """متن خروجی مدل را به (ارزیابی کلی، نقاط قوت، نقاط ضعف) تبدیل می‌کند."""
    text = text or ""
    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    overall = " ".join(_split_items(paragraphs[0])) if paragraphs else ""
    m_s, m_w = "نقاط قوت رزومه", "نقاط ضعف رزومه"
    i_s = text.find(m_s)
    i_w = text.find(m_w, i_s + 1 if i_s != -1 else 0)
    if i_s != -1 and i_w != -1:
        return overall, _split_items(text[i_s + len(m_s):i_w]), _split_items(text[i_w + len(m_w):])
    if i_s != -1:
        return overall, _split_items(text[i_s + len(m_s):]), []
    if i_w != -1:
        return overall, [], _split_items(text[i_w + len(m_w):])
    return " ".join(_split_items(re.sub(r"(?im)^\s*[*_]*(NAME|TECH|EXP|EDU|PROJ)\b.*$", "", text))), [], []


def build_resume_report(ranked_results) -> BytesIO:
    """
    گزارش تحلیلی Word را از روی تمپلیت خام می‌سازد.
    ranked_results: فهرست نتایج به ترتیب رتبه (همان ترتیب فایل Excel).
    هر رزومه دقیقاً از ابتدای یک صفحه‌ی جدید شروع می‌شود (با شکست صفحه‌ی اجباری پیش از آن)، اما دیگر به
    یک صفحه محدود نیست: تحلیل می‌تواند به هر تعداد صفحه که برای پوشش کامل مطالب لازم است ادامه پیدا کند.
    """
    doc = Document(str(REPORT_TEMPLATE_PATH))

    # اگر بدنه‌ی تمپلیت خالی است، پاراگراف‌های خالی را حذف کن تا صفحه‌ی اول خالی نماند
    has_content = bool(doc.tables) or any(
        p.text.strip() or p._p.xpath(".//w:drawing | .//w:pict | .//w:sectPr | .//w:br")
        for p in doc.paragraphs
    )
    if not has_content:
        for p in list(doc.paragraphs):
            p._element.getparent().remove(p._element)

    for idx, r in enumerate(ranked_results, start=1):
        overall, strengths, weaknesses = _parse_resume_analysis(r.get("result_text", ""))
        applicant_line = f"نام متقاضی: {r.get('applicant_name', '-')}"
        name_line = f"نام فایل رزومه: {r['file_name']}"
        level_line = f"سطح ارزیابی: {get_score_level(r['score_raw'])}"

        _add_par(doc, f"رتبه: {idx}", color=REPORT_HEAD_COLOR, page_break=(idx > 1 or has_content))
        _add_par(doc, name_line)
        _add_par(doc, applicant_line)
        _add_par(doc, f"نمره ارزیابی: {r['score_raw']} از 100", color=REPORT_HEAD_COLOR)
        _add_par(doc, level_line, color=REPORT_HEAD_COLOR)

        _add_par(doc, "تحلیل کلی رزومه", color=REPORT_HEAD_COLOR)
        _add_par(doc, overall or "موردی ثبت نشده است.", justify=True)

        _add_par(doc, "نقاط قوت رزومه", color=REPORT_GREEN_COLOR)
        for t in (strengths or ["موردی ثبت نشده است."]):
            _add_par(doc, t, color=REPORT_GREEN_COLOR, bullet=True)

        _add_par(doc, "نقاط ضعف رزومه", color=REPORT_RED_COLOR)
        for t in (weaknesses or ["موردی ثبت نشده است."]):
            _add_par(doc, t, color=REPORT_RED_COLOR, bullet=True)

    out = BytesIO()
    doc.save(out)
    out.seek(0)
    return out


# ==========================================================================================
# Excel Styling - فونت فارسی هم‌فونت گزارش، فونت انگلیسی Cambria
# ==========================================================================================
EXCEL_LATIN_FONT = "Cambria"
EXCEL_FONT_SIZE  = 11
_PERSIAN_CHAR_RE = re.compile(r"[\u0600-\u06FF\u0750-\u077F\uFB50-\uFDFF\uFE70-\uFEFF]")


def _excel_runs(text: str):
    """رشته را به قطعه‌های (متن، فارسی؟) تقسیم می‌کند؛ نویسه‌های خنثی (فاصله، نقطه، خط‌تیره...) به قطعه‌ی مجاور می‌چسبند."""
    kinds = []
    for ch in text:
        if _PERSIAN_CHAR_RE.match(ch):
            kinds.append(True)
        elif ch.isalnum():
            kinds.append(False)
        else:
            kinds.append(None)
    last = None
    for i, k in enumerate(kinds):          # پر کردن خنثی‌ها با نوع قبلی
        if k is None:
            kinds[i] = last
        else:
            last = k
    nxt = None
    for i in range(len(kinds) - 1, -1, -1):  # خنثی‌های ابتدای رشته با نوع بعدی
        if kinds[i] is None:
            kinds[i] = nxt
        else:
            nxt = kinds[i]
    runs = []
    for ch, k in zip(text, kinds):
        k = bool(k)
        if runs and runs[-1][1] == k:
            runs[-1][0] += ch
        else:
            runs.append([ch, k])
    return [(t, k) for t, k in runs]


def style_excel_sheet(ws):
    """
    فونت همه‌ی سلول‌های شیت را تنظیم می‌کند:
    متن فارسی = فونت گزارش تحلیلی (BKoodkBd)، متن و اعداد انگلیسی = Cambria.
    اگر یک سلول ترکیبی باشد (مثلاً نام فایل فارسی + انگلیسی)، هر بخش با فونت خودش نوشته می‌شود.
    """
    family, is_bold = _report_font()
    fa_font = Font(name=family, size=EXCEL_FONT_SIZE, bold=is_bold)
    en_font = Font(name=EXCEL_LATIN_FONT, size=EXCEL_FONT_SIZE)
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if v is None:
                continue
            if not isinstance(v, str):          # اعداد (رتبه، نمره) => Cambria
                cell.font = en_font
                continue
            runs = _excel_runs(v)
            if not runs:
                continue
            kinds = {k for _, k in runs}
            if kinds == {True}:
                cell.font = fa_font
            elif kinds == {False}:
                cell.font = en_font
            else:
                cell.font = en_font
                cell.value = CellRichText(*[
                    TextBlock(
                        InlineFont(rFont=(family if k else EXCEL_LATIN_FONT), sz=EXCEL_FONT_SIZE,
                                   b=(True if (k and is_bold) else None)),
                        t)
                    for t, k in runs
                ])
